# -*- coding: utf-8 -*-
"""OYPB 브랜드별 시딩·원고료·성과 요약표 — 팩트 집계 (추정·안분 없음)"""
import openpyxl
from copy import copy
SRC='/root/.claude/uploads/ea1a909c-9825-524c-af08-aa69686f7bb5/8bba4283-__________.xlsx'
OUT='/home/user/-_oy/reports/OYPB_브랜드별_시딩_비용_광고성과_HAN_260907.xlsx'
wb=openpyxl.load_workbook(SRC); ws=wb['Sheet1']
DEC='_-* #,##0.0_-;\\-* #,##0.0_-;_-* "-"_-;_-@_-'
tpl ={c:{k:copy(getattr(ws.cell(8,c),k)) for k in ('font','fill','border','alignment')} for c in range(2,15)}
tpls={c:{k:copy(getattr(ws.cell(19,c),k)) for k in ('font','fill','border','alignment')} for c in range(2,15)}
nf ={c:ws.cell(8,c).number_format for c in range(2,15)}
nfs={c:ws.cell(19,c).number_format for c in range(2,15)}

# 브랜드 | KR | 글로벌 | 성과집계 | 조회수 | 좋아요 | 댓글 | 원고료 | 광고비 | 매출 | 설명
BD=[
 ('아이디얼포맨',737,0,737,17330588,174558,11787,186700000,33768439,102736712,
  '25년 4Q UGC 협업 첫 시작을 바탕으로 \n250%대 ROAS는 300% 이상으로 올라 섬'),
 ('브링그린',313,0,313,15535558,45278,4314,66300000,124304394,729266288,
  '전체 퍼포먼스 광고비 44% 집행 · 전환 매출 49% 창출\nUGC&협력광고 중심 티트리 라인 스케일업 필요'),
 ('라운드어라운드',281,0,281,2413068,36628,5103,61550000,133848,348140,
  'PR키트 진행 등 참여율 1위(1.73%)\n광고는 281건 중 1건만 집행'),
 ('바이오힐보 (국내)',195,0,195,5440064,22032,2890,37900000,56333229,172219558,
  '퍼포먼스 광고 성과는 전량 국내에서 발생\n최대 조회수 764,503뷰'),
 ('바이오힐보 (해외·JP)',0,7,7,44255,24,5,2331000,0,0,
  '26.06 NAD 크림 나노 시딩 테스트\n원고료 건당 333,000(마크업 포함) · 조회수·참여 저조'),
 ('바이오힐보 (해외·US)',0,68,36,4739117,23068,1066,None,0,0,
  '틱톡 68건 · 청구 30,000,000(2025 시딩 정산 내)\n36건 성과 확보 · 최대 2,954,129뷰 = 전체 1위'),
 ('식물나라',140,0,140,2788821,17873,2446,32490000,14985030,118255596,
  '빅&스몰웨이브의 시작(25.05)\n마이크로 IMC 성과 가장 잘 워킹하는 브랜드'),
 ('웨이크메이크',70,60,109,3143931,17014,1083,40750000,16409851,113788050,
  'JP 60건 건당 415,000 (21건은 26.09 진행 · 미발행)\n26.04 행사·올영픽 미존재로 협업 수량 증대 여지'),
 ('루테카',124,0,124,1655285,8025,1419,22500000,0,0,'퍼포먼스 광고 미집행\n(브랜드 fade out)'),
 ('컬러그램',71,32,103,2653321,8960,1141,31680000,23514230,134432958,
  'IMD 26.09.07로 103건 확정(v8 77건은 부분집합)\nJP 32건 건당 415,000 · ROAS 572%'),
 ('케어플러스',32,0,32,188488,1511,328,3700000,0,0,
  '26.07 UGC 첫 협업 · 퍼포먼스 광고 미집행\n원고료 마이크로 5건×20만 + 나노 27건×10만'),
 ('필리밀리',27,0,27,1353401,1551,416,6300000,14086284,121998372,
  '전체 브랜드 중 ROAS 1위(866%)\n최저 CPA 확보'),
 ('올더베러',8,0,8,396952,3355,85,2000000,0,0,'퍼포먼스 광고 미집행'),
]
r0=8
def style(r,c,val,sub=False):
    cell=ws.cell(r,c); cell.value=val
    src=(tpls if sub else tpl); fmt=(nfs if sub else nf)
    base = src[c] if c<=14 else src[14]
    for k in ('font','fill','border','alignment'): setattr(cell,k,copy(base[k]))
    if c in (11,17): cell.number_format=DEC
    elif c<=14: cell.number_format=fmt[c]
    else: cell.number_format=fmt[6]
for i,(nm,kr,gl,pf,vw,lk,cm,fee,ad,rev,memo) in enumerate(BD):
    r=r0+i
    vals={2:nm,3:kr,4:gl,5:f'=SUM(C{r}:D{r})',6:pf,7:vw,8:lk,9:cm,10:f'=H{r}+I{r}',
          11:f'=J{r}/G{r}',12:(fee if fee else '-'),13:(f'=L{r}/G{r}' if fee else '-'),
          14:ad,15:rev,16:(f'=O{r}/N{r}' if ad else '-'),17:memo}
    for c in range(2,18): style(r,c,vals[c])
    ws.row_dimensions[r].height = 39.6 if '\n' in memo else 26.4
r=r0+len(BD)
S=lambda col: f'=SUM({col}{r0}:{col}{r-1})'
vals={2:'합계',3:S('C'),4:S('D'),5:S('E'),6:S('F'),7:S('G'),8:S('H'),9:S('I'),10:S('J'),
      11:f'=J{r}/G{r}',12:S('L'),13:f'=L{r}/G{r}',14:S('N'),15:S('O'),16:f'=O{r}/N{r}',
      17:'수량 2,165 / 성과 집계 2,112'}
for c in range(2,18): style(r,c,vals[c],sub=True)
ws.row_dimensions[r].height=26.4
sub=r
# 헤더 재구성 — 기존 병합 해제 후 재구성
G6=copy(ws.cell(6,3)._style); G7=copy(ws.cell(7,3)._style)
for mr in list(ws.merged_cells.ranges):
    if mr.min_row in (6,7) or mr.max_row in (6,7): ws.unmerge_cells(str(mr))
for c in range(2,18):
    ws.cell(6,c)._style=copy(G6); ws.cell(7,c)._style=copy(G7)
    ws.cell(6,c).value=None; ws.cell(7,c).value=None
for a,b,val in [(3,6,'UGC 수량'),(7,11,'콘텐츠 성과 (실측)'),(12,13,'원고료'),(14,16,'퍼포먼스 광고')]:
    ws.cell(6,a).value=val; ws.merge_cells(start_row=6,start_column=a,end_row=6,end_column=b)
for c,val in [(2,'브랜드'),(17,'설명')]:
    ws.cell(6,c).value=val; ws.merge_cells(start_row=6,start_column=c,end_row=7,end_column=c)
for c,val in [(3,'KR'),(4,'글로벌'),(5,'전체'),(6,'성과 집계'),(7,'조회수'),(8,'좋아요'),(9,'댓글'),
              (10,'참여'),(11,'ER'),(12,'원고료'),(13,'CPV'),(14,'광고비'),(15,'전환 매출'),(16,'ROAS')]:
    ws.cell(7,c).value=val
for col,w in (('K',9),('L',14),('M',9),('N',14),('O',15),('P',9),('Q',38)): ws.column_dimensions[col].width=w
for rr in range(r0,sub+1): ws.cell(rr,11).number_format='0.00%'
ws['B2']='OYPB 브랜드별 시딩 · 원고료 · 성과  (팩트 집계)'
nfont=copy(ws['B3'].font)
ws['B3']='※ 통계 : 25.04 ~ 26.09 · 국내외 전체 2,165건(KR 1,998 + 글로벌 167 = JP 99 + US 68) · 26.09.07 기준'
ws['B4']='※ RD : 모든 값은 원천 자료 기재값이다. 추정·안분·보정값은 사용하지 않았다.'
for k,t in enumerate([
 '※ 수량 = 글로벌 시딩 라인 자료(JP 99 · US 68) + IMD 집계분 / 성과 집계 = IMD에 성과가 기재된 건수 2,112건 (미집계 53건 = US 32 + JP 26.09 21)',
 '※ 조회수·좋아요·댓글 = IMD 실측값. 좋아요 데이터가 없는 519건(조회수 8,976,727)은 추정하지 않았으므로 참여·ER은 실제보다 낮게(불리하게) 나온 값이다.',
 '※ 원고료 = 국내는 IMD 행 단위 기재값, JP는 확정 단가(웨이크메이크·컬러그램 건당 415,000 / 바이오힐보 건당 333,000 마크업 포함), 케어플러스는 시딩현황 기재값. 바이오힐보 US는 데이터 없음.',
 '※ CPV = 원고료 ÷ 조회수. 브랜드별 세금계산서 귀속은 발행 내역에 브랜드 미귀속액 70,431,189원이 있어 산출하지 않았다.',
 '※ 전체 청구 기준 — 세금계산서 발행 900,344,264원 (26.02~08 확정 840,514,264 + 2025 시딩 정산 59,830,000) ÷ 조회수 57,682,849 = CPV 15.61원 / ÷ 참여 391,960 = CPE 2,297원 / ÷ 2,165건 = 건당 415,863원',
 '※ 2025 시딩은 전량 59,830,000원 내에서 운영 — 바이오힐보 US 틱톡 68건 30,000,000 + 국내 272건 29,830,000',
 '※ JP 미발행 8,715,000원(웨이크메이크 21건 × 415,000) 별도 → 진행 기준 909,059,264원',
]):
    c=ws.cell(sub+1+k,2); c.value=t; c.font=copy(nfont)
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('saved', OUT, '| 합계행', sub)
