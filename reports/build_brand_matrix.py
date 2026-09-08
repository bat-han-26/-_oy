# -*- coding: utf-8 -*-
"""브랜드별 UGC 수량·콘텐츠 성과·퍼포먼스 광고 성과 — 원본 표 레이아웃 1장.
   원천: reports/data.py (SSOT)."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data as D
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

D.verify()
I = D.IDX
B = {b[0]: b for b in D.BRANDS}
BOH = ['바이오힐보 (국내)', '바이오힐보 (해외·JP)', '바이오힐보 (해외·US)']

# label, kind('n' 소계포함 / 's' 참고 소분류), 구성 브랜드, 설명
ROWS = [
 ('아이디얼포맨',      'n', ['아이디얼포맨'],   '25년 4Q UGC 협업 첫 시작 · 최다 협업 브랜드(737건) · ROAS 304%'),
 ('브링그린',          'n', ['브링그린'],       '전체 광고비 43.8% · 전환 매출 48.8% 차지 → 협력광고 중심 스케일업'),
 ('라운드어라운드',    'n', ['라운드어라운드'], '참여율(ER) 1.73% 전체 1위 · 퍼포먼스 광고는 1건만 집행'),
 ('바이오힐보',        'n', BOH,               '국내 195건 + 해외 75건(JP 7 · US 68) 합계 · 해외 성과 집계는 43건분'),
 ('바이오힐보 (국내)', 's', BOH[:1],           '국내만 · 퍼포먼스 광고비·매출은 전액 국내분 (소계 제외)'),
 ('바이오힐보 (해외)', 's', BOH[1:],           'JP 7건(26.06 NAD 크림) + US 68건(25.08 틱톡) · 광고 미집행 (소계 제외)'),
 ('식물나라',          'n', ['식물나라'],       '빅&스몰웨이브의 시작(25.05) · 마이크로 IMC 성과 가장 잘 작동'),
 ('웨이크메이크',      'n', ['웨이크메이크'],   '26.04 행사·올영픽 등 미존재로 UGC 협업 수량 증대해도 좋을 것으로 판단'),
 ('루테카',            'n', ['루테카'],         '퍼포먼스 광고 미집행 (브랜드 fade out)'),
 ('컬러그램',          'n', ['컬러그램'],       'ROAS 572%(브랜드 5위) · JP 32건 포함 → 협업 수량 증가 필요'),
 ('케어플러스',        'n', ['케어플러스'],     '26.07 UGC 첫 협업 · 퍼포먼스 광고 미집행'),
 ('필리밀리',          'n', ['필리밀리'],       '전체 브랜드 중 ROAS 1위(866%) · 최저 CPA 1,129원 확보'),
 ('올더베러',          'n', ['올더베러'],       '퍼포먼스 광고 미집행'),
]

FONT = 'Arial'
def F(sz=9, b=False, color='000000'): return Font(name=FONT, size=sz, bold=b, color=color)
THIN = Side(style='thin', color='BFBFBF'); MED = Side(style='medium', color='808080')
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CTR  = Alignment(horizontal='center', vertical='center', wrap_text=True)
LFT  = Alignment(horizontal='left',   vertical='center', wrap_text=True)
RGT  = Alignment(horizontal='right',  vertical='center')
GRAY = PatternFill('solid', fgColor='D9D9D9'); BLUE = PatternFill('solid', fgColor='DCE6F1')
SUMF = PatternFill('solid', fgColor='F2F2F2')
NUM = '#,##0'; DASH = '#,##0;-#,##0;"-"'

wb = Workbook(); ws = wb.active; ws.title = '브랜드별_성과'
GROUPS = [('브랜드', 1, 1), ('UGC 수량', 2, 4), ('콘텐츠 성과', 5, 9), ('퍼포먼스 광고 성과', 10, 12), ('설명', 13, 13)]
HDR = [None, 'KR', '글로벌', '전체', '조회수', '평균 조회수', '참여수', '평균 CPV', '평균 CPE',
       '광고비', '전환 매출', 'ROAS', None]
for name, c1, c2 in GROUPS:
    ws.merge_cells(start_row=1, start_column=c1, end_row=2 if c1 == c2 else 1, end_column=c2)
    c = ws.cell(1, c1, name); c.font = F(10, True); c.fill = GRAY; c.alignment = CTR
for j, h in enumerate(HDR, 1):
    if h is None: continue
    c = ws.cell(2, j, h); c.font = F(9, True); c.fill = GRAY; c.alignment = CTR
for r in (1, 2):
    for j in range(1, 14):
        ws.cell(r, j).border = Border(left=THIN, right=THIN, top=MED, bottom=MED if r == 2 else THIN)
ws.row_dimensions[1].height = 20; ws.row_dimensions[2].height = 22

g = lambda ms, k: sum(B[m][I[k]] for m in ms)
R0, sub_rows, nn = 3, [], []
for k, (label, kind, ms, memo) in enumerate(ROWS):
    r = R0 + k
    kr, gl, pf = g(ms, 'KR'), g(ms, 'GL'), g(ms, 'PF')
    v, l, c_, cs = g(ms, 'V'), g(ms, 'L'), g(ms, 'C'), g(ms, 'COST')
    ad, rv = g(ms, 'AD'), g(ms, 'REV')
    e = l + c_
    vals = [('  ' if kind == 's' else '') + label, kr or '-', gl or '-', kr + gl, v,
            round(v / pf), e, round(cs / v, 1), round(cs / e), ad or '-', rv or '-']
    for j, x in enumerate(vals, 1): ws.cell(r, j, x)
    ws.cell(r, 12, f'=IF(J{r}="-","-",K{r}/J{r})')
    ws.cell(r, 13, memo)
    (sub_rows if kind == 's' else nn).append(r)

RS = R0 + len(ROWS)
S = lambda col: '=SUM(' + ','.join(f'{col}{r}' for r in nn) + ')'
ws.cell(RS, 1, '소계')
for col in 'BCDEG': ws.cell(RS, ord(col) - 64, S(col))
ws.cell(RS, 6, f'=ROUND(E{RS}/{D.PERF},0)')
ws.cell(RS, 8, f'=ROUND({D.INV_TOT}/E{RS},1)')   # 소계 CPV: 세금계산서 발행 총액 기준
ws.cell(RS, 9, f'=ROUND({D.INV_TOT}/G{RS},0)')   # 소계 CPE: 세금계산서 발행 총액 기준
ws.cell(RS, 10, S('J')); ws.cell(RS, 11, S('K'))
ws.cell(RS, 12, f'=K{RS}/J{RS}'); ws.cell(RS, 13, '-')

for r in range(R0, RS + 1):
    fill = BLUE if r in sub_rows else (SUMF if r == RS else None)
    bold = (r == RS)
    for j in range(1, 14):
        c = ws.cell(r, j); c.font = F(9, bold); c.border = BOX
        if fill: c.fill = fill
        if j == 1: c.alignment = LFT
        elif j == 13: c.alignment = LFT; c.font = F(8, bold, '404040')
        else:
            c.alignment = RGT
            c.number_format = '0%' if j == 12 else ('0.0' if j == 8 else (DASH if j in (2, 3, 10, 11) else NUM))
    ws.row_dimensions[r].height = 26
ws.row_dimensions[RS].height = 20

for w, col in zip([17, 8, 9, 9, 13, 12, 11, 10, 10, 14, 15, 8, 56], 'ABCDEFGHIJKLM'):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B3'
ws.print_area = f'A1:M{RS}'
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

n = RS + 2
for t in [f'※ 수량은 전체 {D.QTY:,}건, 조회수·참여수는 성과 집계 {D.PERF:,}건 기준 (미집계 {D.UNPERF}건 = 바이오힐보 US 32 · 웨이크메이크 JP 26.09 21)',
          f'※ 참여수 = 좋아요+댓글 · 평균 조회수 = 조회수÷성과 집계 건수 · ROAS = 전환 매출÷광고비',
          f'※ 소계의 평균 CPV·CPE는 세금계산서 발행 총액 {D.INV_TOT:,}원 기준 (2026.02~08 확정 {D.INV_2026:,} + 2025 시딩 정산 {D.INV_2025:,})',
          f'※ 브랜드별 평균 CPV·CPE는 브랜드 귀속이 확인된 시딩 비용 {D.COST:,}원 기준(국내는 원고료, 바이오힐보 해외는 청구액) — 발행 총액은 브랜드 귀속 미확정으로 안분하지 않음',
          f'※ 좋아요 공란 {D.LIKE_BLANK}건은 추정하지 않음 → 참여수는 하한, 평균 CPE는 상한. 바이오힐보 (국내)·(해외) 행은 참고용이며 소계에 포함하지 않음']:
    c = ws.cell(n, 1, t); c.font = F(8, color='808080'); c.alignment = LFT
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=13); n += 1

wb.calculation.fullCalcOnLoad = True
OUT = 'reports/OYPB_브랜드별_성과매트릭스_HAN_260908.xlsx'
wb.save(OUT); print('saved', OUT, '· 시트', wb.sheetnames)
